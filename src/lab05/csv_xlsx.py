import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

def csv_to_xlsx(csv_path: str, xlsx_path: str) -> None:
    """Конвертирует CSV в XLSX.
    
    Использует openpyxl.
    Первая строка CSV — заголовок.
    Лист называется "Sheet1".
    Колонки — автоширина по длине текста (не менее 8 символов).
    """
    p_in = Path(csv_path)
    p_out = Path(xlsx_path)

    # Проверка файла
    if not p_in.is_file():
        raise FileNotFoundError(f"Файл не найден: {csv_path}")
    if p_in.suffix.lower() != '.csv':
        raise ValueError("Неверный тип файла, ожидался CSV")

    # Читаем CSV
    with p_in.open("r", newline="", encoding="utf-8") as f:
        reader = csv.reader(f)
        rows = list(reader)

    if not rows:
        raise ValueError("Пустой CSV")

    # Создаем Excel-книгу
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    # Записываем данные и ищем максимальную длину для каждой колонки
    max_lengths = [0] * len(rows[0])

    for r_idx, row in enumerate(rows, start=1):
        for c_idx, value in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c_idx, value=value)
            
            # Запоминаем максимальную длину текста в колонке
            if len(value) > max_lengths[c_idx - 1]:
                max_lengths[c_idx - 1] = len(value)

    # Устанавливаем ширину колонок (не менее 8)
    for i, length in enumerate(max_lengths, start=1):
        col_letter = get_column_letter(i)
        ws.column_dimensions[col_letter].width = max(length + 2, 8)

    # Создаем папки, если их нет, и сохраняем
    p_out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(p_out)

if __name__ == "__main__":
    print("--- ТЕСТ 1: Успешная конвертация в XLSX ---")
    csv_to_xlsx("data/samples/people.csv", "data/out/people.xlsx")
    print("OK: CSV -> XLSX\n")

    print("--- ТЕСТ 2: Ошибка ValueError (неверный формат файла) ---")
    try:
        csv_to_xlsx("data/samples/people.json", "data/out/err.xlsx")
    except ValueError as e:
        print(f"УСПЕШНО ПОЙМАНА ОШИБКА: {e}\n")